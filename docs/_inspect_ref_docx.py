# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt
from pathlib import Path

path = Path(__file__).parent / "all_processes_new.docx"
out = Path(__file__).parent / "_out_inspect.txt"
doc = Document(path)
lines = []
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if not t.strip() and not p.runs:
        lines.append(f"{i:4d} [empty] indent={p.paragraph_format.left_indent}")
        continue
    pf = p.paragraph_format
    ind = pf.left_indent
    fi = pf.first_line_indent
    sn = p.style.name if p.style else "?"
    run_info = []
    for r in p.runs:
        fn = r.font.name
        sz = r.font.size.pt if r.font.size else None
        run_info.append(
            f"'{r.text}' b={r.bold} i={r.italic} f={fn} sz={sz}"
        )
    lines.append(f"{i:4d} [{sn}] ind={ind} fi={fi}")
    lines.append(f"      TEXT: {t!r}")
    for ri in run_info[:8]:
        lines.append(f"      RUN: {ri}")
    if len(run_info) > 8:
        lines.append(f"      ... +{len(run_info)-8} runs")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out, "paras", len(doc.paragraphs))
