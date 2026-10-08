# -*- coding: utf-8 -*-
"""Сборка docs/all_processes.docx по образцу REGISTER sf64avitoP2_список_процессов.docx"""
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph

REGISTER = Path(r"d:\work\REGISTER\2026\0726\300726\sf64avitoP2_список_процессов.docx")
OUT = Path(__file__).resolve().parent / "all_processes.docx"
REGISTER_OUT = REGISTER
BUILD_TMP = Path(__file__).resolve().parent / "_all_processes_built.docx"
TEMPLATE_P1_P6 = Path(__file__).resolve().parent / "sf64avitoP2_template_p1_p6.docx"

TNR = "Times New Roman"
CN = "Courier New"
SZ_BODY = Pt(12)
SZ_CODE = Pt(10)

_CODE_RE = re.compile(
    r"App\\Service\\Avito\\[\w]+|"
    r"Main::[\w]+|"
    r"getStockWithPrice|get_total_kvo_by_profil_store19|"
    r"ExternalFolder::[\w]+|"
    r"invalid access token|refresh token|"
    r"Europe/Moscow|"
    r"sebestoimost_dlya_avito\.xlsx|sebestoimost_\{[^}]+\}\.xlsx|"
    r"[\w_]+(?:\.[\w_]+)+|"
    r"\b(?:avito_product_list|prepared_list_price|prepared_list|products_map|"
    r"inventory_avito|ViewerClaim|MinPrice|total_kvo|pdata|"
    r"AvitoStockApi|AvitoTokenManager|AutoloadClient|PrepareListPrice|"
    r"PrepareListSetKolWithPrice|ExternalFolder|ExcelWriter|file_put_contents|"
    r"delete_prepared_list_price|get_avito_product_list)\b|"
    r"\bTest\d\b|"
    r"[A-Z][a-zA-Z0-9]+(?:[A-Z][a-zA-Z0-9]+)+"
)


def _add_run(p, text: str, *, code: bool = False, bold: bool = False) -> None:
    r = p.add_run(text)
    r.bold = bold
    r.font.name = CN if code else TNR
    r.font.size = SZ_CODE if code else SZ_BODY


def _clear_runs(p: Paragraph) -> None:
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def _is_full_code_line(text: str) -> bool:
    t = text.strip()
    if t.startswith("php bin/console "):
        return True
    if t.startswith(("GET ", "POST ", "PUT ", "SELECT ", "TRUNCATE ")):
        return True
    if t.startswith(("src/", "var/")):
        return True
    if re.match(r"^[\w_]+\s*=\s*.+", t):
        return True
    if t.startswith("Linux:"):
        return True
    if t.startswith("Windows:"):
        return False
    return False


def _fill_service(p: Paragraph, class_path: str) -> None:
    _clear_runs(p)
    _add_run(p, "Команда вызывает сервис ", code=False)
    parts = class_path.replace("App\\", "").split("\\")
    _add_run(p, "App", code=True)
    for part in parts:
        if not part:
            continue
        _add_run(p, "\\", code=True)
        _add_run(p, part, code=True)
    _add_run(p, ".", code=False)


def _fill_mixed(p: Paragraph, text: str) -> None:
    _clear_runs(p)
    if text.startswith("Команда вызывает сервис App\\"):
        fq = text.split("App\\", 1)[1].rstrip(".")
        _fill_service(p, "App\\" + fq)
        return
    if _is_full_code_line(text):
        _add_run(p, text, code=True)
        return
    pos = 0
    for m in _CODE_RE.finditer(text):
        if m.start() > pos:
            _add_run(p, text[pos : m.start()], code=False)
        token = m.group(0)
        if re.fullmatch(r"[а-яА-ЯёЁ]+", token):
            _add_run(p, token, code=False)
        elif token in ("с", "из", "в", "и", "на", "по", "для", "или", "не", "если"):
            _add_run(p, token, code=False)
        else:
            _add_run(p, token, code=True)
        pos = m.end()
    if pos < len(text):
        _add_run(p, text[pos:], code=False)


def _fill_para(p: Paragraph, text: str, style_src: Paragraph) -> None:
    src = style_src.text.strip()
    if style_src.runs and style_src.runs[0].bold and "—" in text:
        _clear_runs(p)
        _add_run(p, text, bold=True)
        return
    if text.startswith("Запускается команда:"):
        _clear_runs(p)
        _add_run(p, text, code=False)
        return
    if text.startswith("php bin/console "):
        _clear_runs(p)
        parts = re.split(r"(bin/console )", text)
        _add_run(p, parts[0], code=True)
        if len(parts) > 2:
            _add_run(p, parts[1], code=True)
            _add_run(p, parts[2], code=True)
        else:
            _add_run(p, text, code=True)
        return
    ppr = style_src._element.pPr
    has_num = ppr is not None and ppr.numPr is not None
    if not has_num and (
        _is_full_code_line(text) or src.startswith(("php ", "GET ", "POST ", "PUT ", "SELECT ", "TRUNCATE ", "src/", "var/"))
        or re.match(r"^[\w_]+\s*=\s*", src)
    ):
        _clear_runs(p)
        if "Команда вызывает" in text:
            _fill_service(p, text.split("сервис ", 1)[1].rstrip("."))
        elif text.startswith("php "):
            _fill_mixed(p, text)
        else:
            _fill_mixed(p, text)
        return
    _fill_mixed(p, text)


def _clone_after(src: Paragraph, after: Paragraph) -> Paragraph:
    new_el = deepcopy(src._element)
    after._element.addnext(new_el)
    return Paragraph(new_el, after._parent)


def _append_pattern(doc: Document, anchor: Paragraph, pattern: list[tuple[int, str]]) -> Paragraph:
    for src_idx, text in pattern:
        src = doc.paragraphs[src_idx]
        new_p = _clone_after(src, anchor)
        _fill_para(new_p, text, src)
        anchor = new_p
    return anchor


def _anchor_after_p6(doc: Document) -> Paragraph:
    for p in reversed(doc.paragraphs):
        if p.text.strip().startswith("Результат P6"):
            return p
    return doc.paragraphs[-1]


def append_p7_p10(doc: Document) -> None:
    anchor = _anchor_after_p6(doc)

    p7 = [
        (39, "P7 — подготовка списка для себестоимости"),
        (40, "Запускается команда:"),
        (41, "php bin/console app:step-prepare-list-price-command"),
        (42, "Команда вызывает сервис App\\Service\\Avito\\PrepareListPrice."),
        (43, "Из avito_product_list выбираются все товары (Main::get_avito_product_list)."),
        (44, "Таблица prepared_list_price очищается (delete_prepared_list_price)."),
        (46, "Для каждого товара создаётся запись:"),
        (47, "product_list_id = avito_product_list.id"),
        (48, "product_id      = avito_product_list.avito_id"),
        (49, "offer_id        = avito_product_list.ad_id"),
        (50, "kol             = 0"),
        (47, "price           = 0"),
        (
            51,
            "Результат P7: staging-таблица prepared_list_price заполнена списком объявлений "
            "для дальнейшего расчёта kol и MinPrice.",
        ),
    ]
    anchor = _append_pattern(doc, anchor, p7)

    p8 = [
        (52, "P8 — kol и MinPrice из ViewerClaim (склад 19)"),
        (53, "Запускается команда:"),
        (54, "php bin/console app:step-prepare-list-set-kol-with-price-command"),
        (55, "Команда вызывает сервис App\\Service\\Avito\\PrepareListSetKolWithPrice."),
        (
            56,
            "Main::PrepareListSetKolWithPrice() читает prepared_list_price, "
            "собирает offer_id как профили.",
        ),
        (
            57,
            "В БД ViewerClaim вызывается getStockWithPrice (остатки и MinPrice, склад 19).",
        ),
        (62, "Для найденных профилей обновляется prepared_list_price:"),
        (63, "kol = total_kvo, price = MinPrice (если задан)."),
        (
            65,
            "Результат P8: в prepared_list_price актуальные kol и себестоимость (price) для экспорта.",
        ),
    ]
    anchor = _append_pattern(doc, anchor, p8)

    p9 = [
        (39, "P9 — экспорт Excel себестоимости"),
        (40, "Запускается команда:"),
        (41, "php bin/console app:step-export-avito-cost-price-excel-command"),
        (43, "Main::ExportAvitoCostPriceExcel() читает prepared_list_price."),
        (44, "ExcelWriter пишет файл:"),
        (45, "src/ModuleAvito/Aion/pdata/sebestoimost_dlya_avito.xlsx"),
        (
            46,
            "Колонки: Номер объявления (product_id), ID объявления (offer_id), "
            "Себестоимость (price).",
        ),
        (51, "Результат P9: xlsx в pdata готов для выгрузки."),
    ]
    anchor = _append_pattern(doc, anchor, p9)

    p10 = [
        (66, "P10 — копия файла во внешнюю папку"),
        (67, "Запускается команда:"),
        (68, "php bin/console app:step-send-avito-cost-price-file-command"),
        (69, "Main::SendAvitoCostPriceFile() читает байты sebestoimost_dlya_avito.xlsx из pdata."),
        (70, "Имя файла: sebestoimost_{dmY_His по Europe/Moscow}.xlsx"),
        (77, "ExternalFolder::save_file() — file_put_contents:"),
        (72, "Windows: {project}/var/data/file/"),
        (73, "Linux: /var/www/txt_files3/"),
        (80, "Результат P10: файл себестоимости доступен внешней системе."),
    ]
    _append_pattern(doc, anchor, p10)


def _copy_built(target: Path, built: Path) -> None:
    import shutil

    try:
        shutil.copy2(built, target)
        print(f"  -> {target}")
    except PermissionError:
        alt = target.with_name(target.stem + "_NEW" + target.suffix)
        shutil.copy2(built, alt)
        print(f"  !! {target} занят — сохранено: {alt}")


def _strip_from_p7(doc: Document) -> None:
    idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("P7"):
            idx = i
            break
    if idx is None:
        return
    for p in list(doc.paragraphs[idx:]):
        el = p._element
        el.getparent().remove(el)


def _ensure_template() -> Path:
    import shutil

    if TEMPLATE_P1_P6.is_file():
        return TEMPLATE_P1_P6
    if not REGISTER.is_file():
        raise FileNotFoundError(f"Нет образца: {REGISTER}")
    doc = Document(REGISTER)
    _strip_from_p7(doc)
    doc.save(TEMPLATE_P1_P6)
    print(f"Сохранён шаблон P1–P6: {TEMPLATE_P1_P6}")
    return TEMPLATE_P1_P6


def build() -> None:
    template = _ensure_template()
    doc = Document(template)
    append_p7_p10(doc)
    out = BUILD_TMP
    try:
        doc.save(out)
    except PermissionError:
        out = BUILD_TMP.with_name("_all_processes_built2.docx")
        doc.save(out)
    print(f"Сборка: {out} (P1–P6 образец + P7–P10 с нумерацией как в P1–P6)")
    built = out
    for target in (OUT, REGISTER_OUT):
        _copy_built(target, built)


if __name__ == "__main__":
    build()
