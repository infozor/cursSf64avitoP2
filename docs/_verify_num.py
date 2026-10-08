from docx import Document
from pathlib import Path

p = Path(__file__).parent / "_all_processes_built2.docx"
d = Document(p)
lines = []
capture = False
for i, para in enumerate(d.paragraphs):
    s = para.text.strip().replace("\n", " ")[:70]
    if s.startswith("P7"):
        capture = True
    if capture and s.startswith("P8"):
        break
    if not capture:
        continue
    num_id = None
    ppr = para._element.pPr
    if ppr is not None and ppr.numPr is not None and ppr.numPr.numId is not None:
        num_id = ppr.numPr.numId.val
    lines.append(f"{i:3d} n={num_id} | {s}")

Path(__file__).parent.joinpath("_verify_p7_num.txt").write_text("\n".join(lines), encoding="utf-8")

# compare P4 snippet
lines2 = []
for i, para in enumerate(d.paragraphs):
    s = para.text.strip().replace("\n", " ")[:70]
    if s.startswith("P4"):
        capture2 = True
    elif s.startswith("P5"):
        break
    else:
        if "capture2" not in dir():
            capture2 = False
        continue
    if not locals().get("capture2"):
        continue
lines2 = []
cap = False
for i, para in enumerate(d.paragraphs):
    s = para.text.strip().replace("\n", " ")[:70]
    if s.startswith("P4 —"):
        cap = True
    if cap and s.startswith("P5 —"):
        break
    if not cap:
        continue
    num_id = None
    ppr = para._element.pPr
    if ppr is not None and ppr.numPr is not None and ppr.numPr.numId is not None:
        num_id = ppr.numPr.numId.val
    lines2.append(f"{i:3d} n={num_id} | {s}")

Path(__file__).parent.joinpath("_verify_p4_num.txt").write_text("\n".join(lines2), encoding="utf-8")
