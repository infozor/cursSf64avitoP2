from docx import Document
from pathlib import Path

p = Path(__file__).parent / "_all_processes_built.docx"
d = Document(p)
headers = []
for x in d.paragraphs:
    s = x.text.strip().replace("\n", " ")
    if s.startswith("P") and len(s) > 2 and s[1].isdigit() and "—" in s:
        headers.append(s)
out = Path(__file__).parent / "_verify.txt"
out.write_text("\n".join(headers) + f"\n\nparagraphs={len(d.paragraphs)}", encoding="utf-8")
print(out.read_text(encoding="utf-8"))
